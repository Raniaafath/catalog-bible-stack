import decimal
import re
import unicodedata
from dataclasses import dataclass
from typing import Dict, List, Optional, Sequence, Tuple

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.models import Q

from catalog.models import (
    Attribute,
    AttributeValue,
    Product,
    ProductAttributeValue,
    ProductType,
    ProductTypeAttribute,
    Variant,
)
from content.models import AttributeI18n, AttributeValueI18n, Locale, ProductTypeI18n, ProductTypeSynonym
from pub.models import Channel


try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - import guard
    raise CommandError("openpyxl is required. Install it before running this command.") from exc


PLACEHOLDER_VALUES = {"non concerné", "non concerne", "n/a", "na", "none", "-", ""}
YES_VALUES = {"oui", "yes", "true", "1"}
NO_VALUES = {"non", "no", "false", "0"}


@dataclass
class ColumnInfo:
    index: int
    code: str
    label: str
    normalized_label: str
    role: str  # sku, product_category, parent_sku, barcode, brand, title, description, media, attribute


@dataclass
class AttributeEntry:
    column: ColumnInfo
    raw_value: object


class Command(BaseCommand):
    help = "Import products + variants + attributes from an Excel export, using only human-readable attribute labels."

    def add_arguments(self, parser):
        parser.add_argument("--path", required=True, help="Path to the Excel file to import.")
        parser.add_argument("--channel", required=True, help="Channel code to scope synonyms and defaults.")
        parser.add_argument(
            "--locale",
            default="fr",
            help="Locale code for labels (Attribute/Value/ProductType i18n). Default: fr.",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Parse and report but roll back database writes.",
        )
        parser.add_argument(
            "--limit",
            type=int,
            default=None,
            help="Only import the first N data rows (after headers). Useful for testing.",
        )

    def handle(self, *args, **options):
        excel_path = options["path"]
        channel_code = options["channel"]
        locale_code = options["locale"]
        dry_run = options["dry_run"]
        limit = options["limit"]

        channel, _ = Channel.objects.get_or_create(code=channel_code, defaults={"name": channel_code})
        locale, _ = Locale.objects.get_or_create(code=locale_code, defaults={"name": locale_code})

        wb = load_workbook(excel_path, read_only=True)
        ws = wb.active

        columns, data_start_row = self._detect_headers(ws)

        metrics = {
            "products_created": 0,
            "products_updated": 0,
            "variants_created": 0,
            "variants_updated": 0,
            "attributes_created": 0,
            "attribute_values_created": 0,
            "placeholders_skipped": 0,
            "rows_processed": 0,
        }

        # Caches to reduce DB lookups.
        attribute_cache: Dict[str, Attribute] = {}
        attribute_value_cache: Dict[Tuple[int, str], AttributeValue] = {}

        with transaction.atomic():
            for row_idx, row in enumerate(
                ws.iter_rows(min_row=data_start_row, values_only=True), start=data_start_row
            ):
                if limit and metrics["rows_processed"] >= limit:
                    break

                row_values = list(row)
                if all(v in (None, "") for v in row_values):
                    continue

                metrics["rows_processed"] += 1
                product_payload, attribute_entries = self._extract_row_payload(
                    columns, row_values, locale, metrics
                )
                if not product_payload.get("product_code"):
                    self.stdout.write(
                        self.style.WARNING(f"Row {row_idx}: skipped (missing product key after cleaning).")
                    )
                    continue

                product_type = self._resolve_product_type(
                    category_label=product_payload.get("category"),
                    locale=locale,
                    channel=channel,
                )

                product, product_created = self._upsert_product(
                    payload=product_payload,
                    product_type=product_type,
                    locale_code=locale.code,
                )
                metrics["products_created"] += int(product_created)
                metrics["products_updated"] += int(not product_created)

                variant, variant_created = self._upsert_variant(product, product_payload)
                metrics["variants_created"] += int(variant_created)
                metrics["variants_updated"] += int(not variant_created)

                self._persist_attributes(
                    product=product,
                    variant=variant,
                    product_type=product_type,
                    attribute_entries=attribute_entries,
                    locale=locale,
                    attribute_cache=attribute_cache,
                    attribute_value_cache=attribute_value_cache,
                    metrics=metrics,
                )

            if dry_run:
                transaction.set_rollback(True)

        self._print_summary(metrics, dry_run)

    # --- Header detection -------------------------------------------------
    def _detect_headers(self, ws) -> Tuple[List[ColumnInfo], int]:
        sampled_rows = list(ws.iter_rows(min_row=1, max_row=10, values_only=True))
        for idx, row in enumerate(sampled_rows):
            normalized = [self._normalize_header_cell(c) for c in row]
            if self._looks_like_code_row(normalized):
                label_row = sampled_rows[idx - 1] if idx > 0 else [""] * len(row)
                columns = self._build_columns(label_row, row)
                return columns, idx + 2  # data starts after the code row (1-based)

            if self._looks_like_label_row(normalized):
                if idx + 1 < len(sampled_rows):
                    candidate_codes = [self._normalize_header_cell(c) for c in sampled_rows[idx + 1]]
                    if self._looks_like_code_row(candidate_codes):
                        columns = self._build_columns(row, sampled_rows[idx + 1])
                        return columns, idx + 3

        # Fallback: pick the first row containing shop_sku as codes, with previous row as labels.
        for idx, row in enumerate(sampled_rows):
            raw_cells = [str(c).strip().lower() if c is not None else "" for c in row]
            if any(cell in {"shop_sku", "shop sku", "sku"} for cell in raw_cells):
                label_row = sampled_rows[idx - 1] if idx > 0 else [""] * len(row)
                columns = self._build_columns(label_row, row)
                return columns, idx + 2

        raise CommandError("Could not locate header rows (expected labels + codes with a shop_sku column).")

    def _looks_like_code_row(self, normalized_cells: Sequence[str]) -> bool:
        has_sku = any(cell in {"shop_sku", "sku"} for cell in normalized_cells)
        feature_count = sum(1 for cell in normalized_cells if cell.startswith("feature_") or cell.startswith("att_"))
        return has_sku or feature_count > 5

    def _looks_like_label_row(self, normalized_cells: Sequence[str]) -> bool:
        return any("identifiant" in cell or "sku" in cell for cell in normalized_cells)

    def _build_columns(self, label_row: Sequence[object], code_row: Sequence[object]) -> List[ColumnInfo]:
        columns: List[ColumnInfo] = []
        for idx, (label_raw, code_raw) in enumerate(zip(label_row, code_row)):
            label_clean, normalized_label = self._normalize_label(label_raw or code_raw or "")
            code_str = str(code_raw or "").strip()
            role = self._detect_role(label_clean, normalized_label, code_str)
            columns.append(
                ColumnInfo(
                    index=idx,
                    code=code_str,
                    label=label_clean,
                    normalized_label=normalized_label,
                    role=role,
                )
            )
        return columns

    def _detect_role(self, label: str, normalized_label: str, code: str) -> str:
        code_lower = code.lower()
        if code_lower in {"product_category"} or "categorie" in normalized_label or "category" in normalized_label:
            return "product_category"
        if code_lower in {"shop_sku", "sku"} or "identifiant-produit" in normalized_label or "sku" in normalized_label:
            return "sku"
        if "parent" in code_lower or "parent" in normalized_label:
            return "parent_sku"
        if "ean" in code_lower or "gtin" in code_lower or "ean" in normalized_label or "barcode" in normalized_label:
            return "barcode"
        if "brand" in code_lower or "marque" in normalized_label:
            return "brand"
        if "title" in code_lower or "titre" in normalized_label:
            return "title"
        if "description" in code_lower:
            return "description"
        if code_lower.startswith("media_") or "image" in normalized_label or "pdf" in normalized_label:
            return "media"
        if code_lower in {"product_administrative_code", "reference"}:
            return "internal_code"
        return "attribute"

    def _normalize_header_cell(self, cell: object) -> str:
        text, normalized = self._normalize_label(cell or "")
        return normalized

    # --- Row parsing ------------------------------------------------------
    def _extract_row_payload(
        self,
        columns: Sequence[ColumnInfo],
        row_values: Sequence[object],
        locale: Locale,
        metrics: Dict[str, int],
    ) -> Tuple[Dict[str, object], List[AttributeEntry]]:
        payload: Dict[str, object] = {}
        attributes: List[AttributeEntry] = []

        for col_meta, raw_value in zip(columns, row_values):
            if raw_value in (None, ""):
                continue

            cleaned_value = self._clean_value(raw_value)
            if cleaned_value is None:
                metrics["placeholders_skipped"] += 1
                continue

            if col_meta.role == "product_category":
                payload["category"] = cleaned_value
            elif col_meta.role == "sku":
                payload["sku"] = cleaned_value
            elif col_meta.role == "parent_sku":
                payload["parent_sku"] = cleaned_value
            elif col_meta.role == "barcode":
                payload["barcode"] = cleaned_value
            elif col_meta.role == "brand":
                payload["brand"] = cleaned_value
            elif col_meta.role == "title":
                payload["title"] = cleaned_value
            elif col_meta.role == "description":
                payload["description"] = cleaned_value
            elif col_meta.role == "internal_code":
                payload["internal_code"] = cleaned_value
            elif col_meta.role == "media":
                # Media is acknowledged but not imported in this command.
                continue
            else:
                attributes.append(AttributeEntry(column=col_meta, raw_value=cleaned_value))

        product_code = self._choose_product_code(payload)
        variant_code = payload.get("sku") or payload.get("internal_code")
        payload["product_code"] = product_code
        payload["variant_code"] = variant_code
        return payload, attributes

    def _choose_product_code(self, payload: Dict[str, object]) -> Optional[str]:
        sku = payload.get("sku")
        if sku:
            return sku
        parent = payload.get("parent_sku")
        if parent:
            return parent
        brand = payload.get("brand")
        internal = payload.get("internal_code")
        if brand and internal:
            return f"{brand}-{internal}"
        if internal:
            return internal
        title = payload.get("title")
        if brand and title:
            return f"{brand}-{self._slugify(title)[:60]}"
        return None

    def _clean_value(self, raw_value: object) -> Optional[object]:
        if raw_value is None:
            return None
        if isinstance(raw_value, float):
            if raw_value != raw_value:  # NaN
                return None
        text = str(raw_value).strip()
        text_compact = re.sub(r"\s+", " ", text)
        if self._is_placeholder(text_compact):
            return None
        return text_compact

    # --- Upserts ----------------------------------------------------------
    def _resolve_product_type(self, category_label: Optional[str], locale: Locale, channel: Channel) -> ProductType:
        label = self._normalize_category_label(category_label or "Divers")
        label_clean, normalized_label = self._normalize_label(label)

        product_type = None

        synonym = (
            ProductTypeSynonym.objects.filter(
                locale=locale,
                term__iexact=label_clean,
            )
            .filter(Q(channel=channel) | Q(channel__isnull=True))
            .order_by("-is_active", "-priority")
            .first()
        )
        if synonym:
            product_type = synonym.product_type

        if not product_type:
            pti18n = (
                ProductTypeI18n.objects.select_related("product_type")
                .filter(locale=locale, label__iexact=label_clean)
                .first()
            )
            if pti18n:
                product_type = pti18n.product_type

        if not product_type:
            code = self._unique_code(ProductType, normalized_label or "type")
            product_type, created = ProductType.objects.get_or_create(
                code=code,
                defaults={
                    "default_label": label_clean,
                    "category_path": label_clean,
                },
            )
            if created:
                self.stdout.write(self.style.SUCCESS(f"Created product type {product_type.code} ({label_clean})"))

        ProductTypeI18n.objects.get_or_create(
            product_type=product_type,
            locale=locale,
            defaults={"label": label_clean},
        )
        return product_type

    def _normalize_category_label(self, label: str) -> str:
        if not label:
            return label
        # Keep the last path segment and normalize whitespace/accents.
        parts = [part.strip() for part in str(label).split("/") if part and str(part).strip()]
        segment = parts[-1] if parts else str(label)
        normalized = unicodedata.normalize("NFC", segment)
        normalized = normalized.replace("\u00a0", " ")
        normalized = re.sub(r"\s+", " ", normalized).strip()
        # Tighten spacing around apostrophes/hyphens.
        normalized = re.sub(r"\s*([’'‑-])\s*", r"\1", normalized)
        return normalized

    def _upsert_product(
        self, payload: Dict[str, object], product_type: ProductType, locale_code: str
    ) -> Tuple[Product, bool]:
        defaults = {
            "product_type": product_type,
            "brand": payload.get("brand") or "",
            "default_label": payload.get("title") or "",
            "source_title": payload.get("title") or "",
            "source_description": payload.get("description") or "",
            "source_locale": locale_code,
            "source_supplier": payload.get("brand") or "",
            "source_sku": payload.get("sku") or "",
        }
        product, created = Product.objects.get_or_create(code=payload["product_code"], defaults=defaults)
        updates = {}
        if product.product_type_id != product_type.id:
            product.product_type = product_type
            updates["product_type"] = product_type
        for field in ("brand", "default_label", "source_title", "source_description", "source_sku"):
            desired = defaults[field]
            if getattr(product, field) != desired and desired:
                setattr(product, field, desired)
                updates[field] = desired
        if updates:
            product.save(update_fields=list(updates.keys()))
        return product, created

    def _upsert_variant(self, product: Product, payload: Dict[str, object]) -> Tuple[Variant, bool]:
        sku = payload.get("variant_code")
        variant = None
        created = False
        if sku:
            variant = Variant.objects.filter(sku=sku).first()
        if variant:
            updates = {}
            if variant.product_id != product.id:
                variant.product = product
                updates["product"] = product
            barcode = payload.get("barcode")
            if barcode and variant.barcode != barcode:
                variant.barcode = barcode
                updates["barcode"] = barcode
            if updates:
                variant.save(update_fields=list(updates.keys()))
            return variant, False

        variant = Variant(product=product, sku=sku, barcode=payload.get("barcode"))
        variant.save()
        created = True
        return variant, created

    # --- Attribute persistence -------------------------------------------
    def _persist_attributes(
        self,
        product: Product,
        variant: Variant,
        product_type: ProductType,
        attribute_entries: Sequence[AttributeEntry],
        locale: Locale,
        attribute_cache: Dict[str, Attribute],
        attribute_value_cache: Dict[Tuple[int, str], AttributeValue],
        metrics: Dict[str, int],
    ) -> None:
        # Dimension handling: collect values up front.
        dimension_buffer: Dict[str, AttributeEntry] = {}
        other_entries: List[AttributeEntry] = []
        for entry in attribute_entries:
            if self._is_dimension_label(entry.column.normalized_label):
                dimension_buffer[entry.column.normalized_label] = entry
            else:
                other_entries.append(entry)

        dimension_entries = self._prepare_dimension_entries(dimension_buffer)
        all_entries = dimension_entries + other_entries

        variant_signature_parts: List[str] = []

        for entry in all_entries:
            values = self._split_multi_values(entry.raw_value)
            if not values:
                metrics["placeholders_skipped"] += 1
                continue
            is_multi = len(values) > 1

            # Decide scope: variant-level for color/size-like attributes when we have a variant code.
            variant_level = bool(
                variant.sku
                and (
                    "couleur" in entry.column.normalized_label
                    or "color" in entry.column.normalized_label
                    or "taille" in entry.column.normalized_label
                    or "format" in entry.column.normalized_label
                    or "dimension" in entry.column.normalized_label
                )
            )

            attribute = self._get_or_create_attribute(
                label=entry.column.label,
                normalized_label=entry.column.normalized_label,
                sample_value=values[0],
                is_multi=is_multi,
                variant_level=variant_level,
                product_type=product_type,
                locale=locale,
                cache=attribute_cache,
                metrics=metrics,
            )

            if variant_level:
                ProductTypeAttribute.objects.update_or_create(
                    product_type=product_type,
                    attribute=attribute,
                    defaults={"variant_level": True},
                )

            value_payloads = [self._normalize_value(v) for v in values]
            value_payloads = [v for v in value_payloads if v is not None]
            # Dedupe repeated values after normalization to keep idempotent writes.
            deduped: List[Dict[str, object]] = []
            seen_labels = set()
            for payload in value_payloads:
                label_key = str(payload.get("label", "")).lower()
                if label_key in seen_labels:
                    continue
                seen_labels.add(label_key)
                deduped.append(payload)
            value_payloads = deduped
            if not value_payloads:
                metrics["placeholders_skipped"] += 1
                continue

            pav_scope = {"variant": variant} if variant_level else {"product": product}

            # Handle multi-value by storing JSON of labels/codes, while keeping the first value in attribute_value/value_number for compatibility.
            first_value = value_payloads[0]
            attr_value_obj = None
            value_number = None
            value_text = None
            value_json = None
            unit = first_value.get("unit")

            if first_value["kind"] == "enum":
                attr_value_obj = self._get_or_create_attribute_value(
                    attribute,
                    first_value["label"],
                    locale,
                    attribute_value_cache,
                    metrics,
                )
            elif first_value["kind"] == "number":
                value_number = first_value["number"]
                value_text = first_value["label"]
                unit = first_value.get("unit")
            else:
                value_text = first_value["label"]

            if len(value_payloads) > 1:
                serialized = []
                for payload in value_payloads:
                    if payload["kind"] == "enum":
                        val_obj = self._get_or_create_attribute_value(
                            attribute,
                            payload["label"],
                            locale,
                            attribute_value_cache,
                            metrics,
                        )
                        serialized.append({"code": val_obj.code, "label": payload["label"]})
                    else:
                        serialized.append({"label": payload["label"], "unit": payload.get("unit")})
                value_json = serialized
                attribute.is_multi = True
                attribute.save(update_fields=["is_multi"])

            pav_defaults = {
                "attribute_value": attr_value_obj,
                "value_number": value_number,
                "value_text": value_text,
                "value_json": value_json,
                "unit": unit,
            }
            pav, created = ProductAttributeValue.objects.get_or_create(
                attribute=attribute,
                defaults={**pav_scope, **pav_defaults},
                **pav_scope,
            )
            if not created:
                updates = {}
                if pav.attribute_value_id != (attr_value_obj.id if attr_value_obj else None):
                    pav.attribute_value = attr_value_obj
                    updates["attribute_value"] = attr_value_obj
                if pav.value_number != value_number:
                    pav.value_number = value_number
                    updates["value_number"] = value_number
                if pav.value_text != value_text:
                    pav.value_text = value_text
                    updates["value_text"] = value_text
                if pav.value_json != value_json:
                    pav.value_json = value_json
                    updates["value_json"] = value_json
                if unit and pav.unit != unit:
                    pav.unit = unit
                    updates["unit"] = unit
                if updates:
                    pav.save(update_fields=list(updates.keys()))

            if variant_level:
                signature_fragment = self._build_signature_fragment(attribute, value_payloads)
                if signature_fragment:
                    variant_signature_parts.append(signature_fragment)

        if variant_signature_parts:
            signature = "|".join(sorted(variant_signature_parts))
            if variant.axis_signature != signature:
                variant.axis_signature = signature
                variant.save(update_fields=["axis_signature"])

    def _build_signature_fragment(self, attribute: Attribute, values: Sequence[Dict[str, object]]) -> str:
        labels = [v["label"] for v in values if v.get("label")]
        if not labels:
            return ""
        return f"{attribute.code}:{'/'.join(labels)}"

    def _get_or_create_attribute(
        self,
        label: str,
        normalized_label: str,
        sample_value: object,
        is_multi: bool,
        variant_level: bool,
        product_type: ProductType,
        locale: Locale,
        cache: Dict[str, Attribute],
        metrics: Dict[str, int],
    ) -> Attribute:
        cache_key = normalized_label or self._normalize_label(label)[1] or label.lower()
        if cache_key in cache:
            attr_cached = cache[cache_key]
            self._ensure_attribute_links(
                attribute=attr_cached,
                product_type=product_type,
                locale=locale,
                label=label,
                variant_level=variant_level,
            )
            return attr_cached

        data_type, unit = self._infer_attribute_type(sample_value)

        attr = None
        if normalized_label:
            slug = self._slugify(normalized_label)
            attr = Attribute.objects.filter(code=slug).first()
        if not attr:
            attr = Attribute.objects.filter(i18n__label__iexact=label, i18n__locale=locale).first()

        if not attr:
            code = self._unique_code(Attribute, normalized_label or label or "attribute")
            attr = Attribute.objects.create(
                code=code,
                data_type=data_type,
                unit=unit,
                is_multi=is_multi,
            )
            metrics["attributes_created"] += 1
            cache[cache_key] = attr
            self._ensure_attribute_links(
                attribute=attr,
                product_type=product_type,
                locale=locale,
                label=label,
                variant_level=variant_level,
            )
            return attr

        updates = {}
        if unit and not attr.unit:
            attr.unit = unit
            updates["unit"] = unit
        if is_multi and not attr.is_multi:
            attr.is_multi = True
            updates["is_multi"] = True
        if updates:
            attr.save(update_fields=list(updates.keys()))

        self._ensure_attribute_links(
            attribute=attr,
            product_type=product_type,
            locale=locale,
            label=label,
            variant_level=variant_level,
        )
        cache[cache_key] = attr
        return attr

    def _ensure_attribute_links(
        self,
        attribute: Attribute,
        product_type: ProductType,
        locale: Locale,
        label: str,
        variant_level: bool,
    ) -> None:
        pta, created = ProductTypeAttribute.objects.get_or_create(
            product_type=product_type,
            attribute=attribute,
            defaults={"variant_level": variant_level},
        )
        if not created and variant_level and not pta.variant_level:
            pta.variant_level = True
            pta.save(update_fields=["variant_level"])
        AttributeI18n.objects.get_or_create(attribute=attribute, locale=locale, defaults={"label": label})

    def _get_or_create_attribute_value(
        self,
        attribute: Attribute,
        label: str,
        locale: Locale,
        cache: Dict[Tuple[int, str], AttributeValue],
        metrics: Dict[str, int],
    ) -> AttributeValue:
        _, normalized = self._normalize_label(label)
        cache_key = (attribute.id, normalized)
        if cache_key in cache:
            return cache[cache_key]

        code = self._slugify(normalized or label)
        value_obj, created = AttributeValue.objects.get_or_create(attribute=attribute, code=code)
        if created:
            metrics["attribute_values_created"] += 1
        AttributeValueI18n.objects.get_or_create(
            attribute_value=value_obj,
            locale=locale,
            defaults={"label": label},
        )
        cache[cache_key] = value_obj
        return value_obj

    # --- Helpers ----------------------------------------------------------
    def _normalize_value(self, raw: object) -> Optional[Dict[str, object]]:
        if raw is None:
            return None
        label, normalized = self._normalize_label(raw)
        if not label:
            return None

        lowered = label.lower()
        if lowered in YES_VALUES:
            return {"kind": "enum", "label": "Oui"}
        if lowered in NO_VALUES:
            return {"kind": "enum", "label": "Non"}

        number_match = re.match(r"^(-?\\d+[\\.,]?\\d*)(?:\\s*([a-zA-Z%°]+))?$", label)
        if number_match:
            number_raw = number_match.group(1).replace(",", ".")
            try:
                number_val = decimal.Decimal(number_raw)
                unit = number_match.group(2)
                return {"kind": "number", "label": label, "number": number_val, "unit": unit}
            except decimal.InvalidOperation:
                pass

        if len(label) > 150:
            return {"kind": "text", "label": label}

        return {"kind": "enum", "label": label}

    def _infer_attribute_type(self, sample_value: object) -> Tuple[str, Optional[str]]:
        normalized = self._normalize_value(sample_value)
        if not normalized:
            return Attribute.DataType.TEXT, None
        if normalized["kind"] == "number":
            return Attribute.DataType.NUMBER, normalized.get("unit")
        if normalized["kind"] == "text":
            return Attribute.DataType.TEXT, None
        return Attribute.DataType.ENUM, None

    def _split_multi_values(self, raw_value: object) -> List[object]:
        if raw_value is None:
            return []
        if isinstance(raw_value, (int, float, decimal.Decimal)):
            return [raw_value]
        text = str(raw_value)
        if ";" in text or "|" in text:
            parts = re.split(r"[;|]", text)
        elif "," in text and not re.match(r"^\\s*\\d+,\\d+\\s*$", text):
            parts = text.split(",")
        else:
            parts = [text]
        cleaned_parts = []
        for part in parts:
            cleaned = part.strip()
            if not cleaned or self._is_placeholder(cleaned):
                continue
            cleaned_parts.append(cleaned)
        return cleaned_parts

    def _is_dimension_label(self, normalized_label: str) -> bool:
        return any(
            token in normalized_label
            for token in [
                "longueur",
                "largeur",
                "hauteur",
                "profondeur",
                "dimensions",
                "dimension",
                "epaisseur",
            ]
        )

    def _prepare_dimension_entries(self, entries: Dict[str, AttributeEntry]) -> List[AttributeEntry]:
        if not entries:
            return []

        # Check existing schema to decide whether to split or keep combined.
        separate_exists = Attribute.objects.filter(
            Q(code__in=["longueur", "largeur", "hauteur", "profondeur", "epaisseur"])
            | Q(i18n__label__in=["Longueur", "Largeur", "Hauteur", "Profondeur", "Épaisseur"])
        ).exists()
        has_dimension_col = any("dimension" in key for key in entries.keys())

        if separate_exists:
            return list(entries.values())

        if has_dimension_col:
            # Prefer the provided combined column if schema does not enforce separate attributes.
            first_key = next(iter(entries.keys()))
            return [entries[first_key]]

        # Build a combined "Dimensions" attribute from individual parts.
        def find_value(token: str) -> Optional[AttributeEntry]:
            for key, entry in entries.items():
                if token in key:
                    return entry
            return None

        parts = []
        for prefix, token in (("L", "longueur"), ("l", "largeur"), ("H", "hauteur"), ("P", "profondeur")):
            entry = find_value(token)
            if entry:
                parts.append(f"{prefix} {entry.raw_value}")

        if parts:
            synthetic_col = ColumnInfo(
                index=-1,
                code="dimensions",
                label="Dimensions",
                normalized_label="dimensions",
                role="attribute",
            )
            return [AttributeEntry(column=synthetic_col, raw_value=" x ".join(parts))]

        return list(entries.values())

    def _normalize_label(self, text: object) -> Tuple[str, str]:
        if text is None:
            return "", ""
        raw = str(text)
        collapsed = re.sub(r"\\s+", " ", raw).strip()
        cleaned = re.sub(r"[\\s\\u00A0]+", " ", collapsed)
        normalized = re.sub(r"[^\\w]+", " ", cleaned, flags=re.UNICODE).strip().lower()
        normalized = re.sub(r"\\s+", "-", normalized).strip("-")
        return cleaned, normalized

    def _slugify(self, text: str) -> str:
        _, normalized = self._normalize_label(text)
        return normalized or "value"

    def _unique_code(self, model, base: str, field: str = "code") -> str:
        slug = self._slugify(base)
        candidate = slug or "item"
        suffix = 2
        while model.objects.filter(**{field: candidate}).exists():
            candidate = f"{slug}-{suffix}"
            suffix += 1
        return candidate

    def _is_placeholder(self, text: str) -> bool:
        return text.lower().strip() in PLACEHOLDER_VALUES

    def _print_summary(self, metrics: Dict[str, int], dry_run: bool) -> None:
        prefix = "[DRY-RUN] " if dry_run else ""
        self.stdout.write(
            f"{prefix}Processed {metrics['rows_processed']} rows | "
            f"products created/updated: {metrics['products_created']}/{metrics['products_updated']} | "
            f"variants created/updated: {metrics['variants_created']}/{metrics['variants_updated']} | "
            f"attributes created: {metrics['attributes_created']} | "
            f"attribute values created: {metrics['attribute_values_created']} | "
            f"placeholders skipped: {metrics['placeholders_skipped']}"
        )
