import hashlib
import os
import re
from typing import Dict, List, Optional, Tuple

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from catalog.models import (
    Attribute,
    AttributeValue,
    Product,
    ProductAttributeValue,
    ProductType,
    ProductTypeAttribute,
    Variant,
)
from content.models import AttributeI18n, AttributeValueI18n, Locale
from pub.models import Channel, Template, TemplatePart


class Command(BaseCommand):
    help = "Seed FR locale/channel/product plus attributes from the first non-empty row of the Excel export, and ensure a simple FR title template."

    CODE_OVERRIDES: Dict[str, str] = {
        "feature_00562_201010|receveur_de_doucher_a_encastrer|receveur_de_douche|r07-001-005": "forme",
        "feature_01758_201010|receveur_de_doucher_a_encastrer|receveur_de_douche|r07-001-005": "type-de-pose",
        "feature_22088_201010|receveur_de_doucher_a_encastrer|receveur_de_douche|r07-001-005": "type-de-produit",
        "feature_10779_201010|receveur_de_doucher_a_encastrer|receveur_de_douche|r07-001-005": "famille-de-matiere",
        "feature_08547_201010|receveur_de_doucher_a_encastrer|receveur_de_douche|r07-001-005": "nom-produit",
        "feature_01181_201010|receveur_de_doucher_a_encastrer|receveur_de_douche|r07-001-005": "type-produit-alt",
        "feature_15092_201010|receveur_de_doucher_a_encastrer|receveur_de_douche|r07-001-005": "glissance-pn",
        "feature_26094_201010|receveur_de_doucher_a_encastrer|receveur_de_douche|r07-001-005": "glissance-une-41901",
        "feature_04861_201010|receveur_de_doucher_a_encastrer|receveur_de_douche|r07-001-005": "emplacement-vidage",
        "feature_03108_201010|receveur_de_doucher_a_encastrer|receveur_de_douche|r07-001-005": "recoupable",
        "feature_14428_201010|receveur_de_doucher_a_encastrer|receveur_de_douche|r07-001-005": "effet",
        "feature_10847_201010|receveur_de_doucher_a_encastrer|receveur_de_douche|r07-001-005": "portrait-client",
        "feature_10842_201010|receveur_de_doucher_a_encastrer|receveur_de_douche|r07-001-005": "finition-produit",
        "feature_06989_201010|receveur_de_doucher_a_encastrer|receveur_de_douche|r07-001-005": "accessoires-fournis",
        "feature_15095_201010|receveur_de_doucher_a_encastrer|receveur_de_douche|r07-001-005": "bonde-vidage",
    }

    def add_arguments(self, parser):
        parser.add_argument(
            "--excel",
            default="excel/export-products-20251219134523.xlsx",
            help="Path to the Excel export to read sample data from.",
        )
        parser.add_argument(
            "--channel",
            default="default",
            help="Channel code to seed (will be created if missing).",
        )
        parser.add_argument(
            "--product-type",
            default=None,
            help="Override product_type code. Defaults to a slug of the product_category cell.",
        )

    def handle(self, *args, **options):
        excel_path = options["excel"]
        channel_code = options["channel"]
        product_type_override = options.get("product_type")

        sample_row = self._load_sample_row(excel_path)
        if not sample_row:
            raise CommandError("No usable row found in the Excel file.")

        product_category = sample_row["product_category"]
        shop_sku = sample_row["shop_sku"]
        brand = sample_row["brand"]
        title_fr = sample_row["title_fr"]
        description_fr = sample_row["description_fr"]
        attr_rows = sample_row["attributes"]

        product_type_code = product_type_override or self._slugify(product_category.split("/")[-1])

        with transaction.atomic():
            locale_fr, _ = Locale.objects.get_or_create(code="fr", defaults={"name": "Français"})
            channel, _ = Channel.objects.get_or_create(code=channel_code, defaults={"name": channel_code})
            product_type, _ = ProductType.objects.get_or_create(
                code=product_type_code, defaults={"default_label": product_category}
            )

            product, _ = Product.objects.get_or_create(
                code=shop_sku,
                defaults={
                    "product_type": product_type,
                    "brand": brand,
                    "default_label": title_fr,
                    "source_title": title_fr or "",
                    "source_description": description_fr or "",
                    "source_locale": "fr",
                    "source_supplier": "excel_seed",
                    "source_sku": shop_sku,
                },
            )
            variant, _ = Variant.objects.get_or_create(product=product, sku=shop_sku)

            attr_by_code: Dict[str, Attribute] = {}
            for attr_code, attr_label_fr, value in attr_rows:
                if attr_code in ("product_category", "shop_sku"):
                    continue
                attr = self._ensure_attribute(attr_code, attr_label_fr, value, locale_fr, product_type)
                attr_by_code[attr_code] = attr
                if self._is_numeric(value):
                    self._upsert_pav(product, attr, value_text=str(value))
                else:
                    val_obj = self._ensure_enum_value(attr, str(value), locale_fr)
                    self._upsert_pav(product, attr, attr_value=val_obj)

            template = (
                Template.objects.filter(
                    product_type=product_type,
                    locale=locale_fr,
                    channel=channel,
                    kind=Template.Kind.TITLE,
                    status=Template.Status.ACTIVE,
                )
                .order_by("-version", "-id")
                .first()
            )
            if not template:
                template = Template.objects.create(
                    product_type=product_type,
                    locale=locale_fr,
                    channel=channel,
                    kind=Template.Kind.TITLE,
                    version=1,
                    status=Template.Status.ACTIVE,
                )
                self.stdout.write(self.style.SUCCESS(f"Created template {template}"))

            template_parts: List[Tuple[int, TemplatePart.PartType, Optional[Attribute], Optional[str]]] = []
            template_parts.append((1, TemplatePart.PartType.BRAND, None, None))
            for code in [
                "feature_22088_201010|RECEVEUR_DE_DOUCHER_A_ENCASTRER|RECEVEUR_DE_DOUCHE|R07-001-005",  # type de produit
                "feature_01758_201010|RECEVEUR_DE_DOUCHER_A_ENCASTRER|RECEVEUR_DE_DOUCHE|R07-001-005",  # type de pose
                "feature_00562_201010|RECEVEUR_DE_DOUCHER_A_ENCASTRER|RECEVEUR_DE_DOUCHE|R07-001-005",  # forme
                "feature_10837_main_color",  # couleur
                "feature_10840_main_material",  # matière
                "ATT_15836",  # dimensions
                "feature_14428_201010|RECEVEUR_DE_DOUCHER_A_ENCASTRER|RECEVEUR_DE_DOUCHE|R07-001-005",  # effet
            ]:
                attr = attr_by_code.get(code)
                if attr:
                    template_parts.append((len(template_parts) + 1, TemplatePart.PartType.ATTRIBUTE_VALUE, attr, None))

            if len(template_parts) == 1:
                self.stdout.write(self.style.WARNING("No attributes found for template parts; template only has brand."))

            parts = template_parts
            for position, part_type, attribute, literal_text in parts:
                part, created = TemplatePart.objects.get_or_create(
                    template=template,
                    position=position,
                    defaults={
                        "part_type": part_type,
                        "attribute": attribute,
                        "literal_text": literal_text,
                        "required": False,
                    },
                )
                if not created:
                    update_fields = {}
                    if part.part_type != part_type:
                        part.part_type = part_type
                        update_fields["part_type"] = part_type
                    if part.attribute_id != (attribute.id if attribute else None):
                        part.attribute = attribute
                        update_fields["attribute"] = attribute
                    if part.literal_text != literal_text:
                        part.literal_text = literal_text
                        update_fields["literal_text"] = literal_text
                    if update_fields:
                        part.save(update_fields=list(update_fields.keys()))
                action = "Created" if created else "Ensured"
                self.stdout.write(self.style.SUCCESS(f"{action} part #{position} ({part.part_type})"))

                self.stdout.write(
                    self.style.SUCCESS(
                        f"Seeded product {product.code} / variant {variant.id} with FR template {template.id}."
                    )
                )

    def _load_sample_row(self, excel_path: str) -> Optional[Dict[str, object]]:
        if not os.path.exists(excel_path):
            raise CommandError(f"Excel file not found: {excel_path}")
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError("openpyxl is required. Install with `pip install openpyxl`.") from exc

        wb = load_workbook(excel_path)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 3:
            return None
        headers_fr = rows[0]
        headers_codes = rows[1]

        for row in rows[2:]:
            row_map = dict(zip(headers_codes, row))
            sku = row_map.get("shop_sku")
            if not sku:
                continue
            attrs = []
            for fr_label, code, val in zip(headers_fr, headers_codes, row):
                if val in (None, ""):
                    continue
                if not (code.startswith("feature_") or code.startswith("ATT_")):
                    continue
                attrs.append((code, fr_label or code, val))

            return {
                "product_category": row_map.get("product_category") or "",
                "shop_sku": sku,
                "brand": row_map.get("feature_06575_brand") or "",
                "title_fr": row_map.get("i18n_fr_12963_title") or "",
                "description_fr": row_map.get("i18n_fr_01022_longdescription") or "",
                "attributes": attrs,
            }
        return None

    def _slugify(self, text: str) -> str:
        slug = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
        return slug or "product-type"

    def _ensure_attribute(
        self, code: str, fr_label: str, sample_value: object, locale: Locale, product_type: ProductType
    ) -> Attribute:
        is_numeric = self._is_numeric(sample_value)
        data_type = Attribute.DataType.TEXT if is_numeric else Attribute.DataType.ENUM
        override = self.CODE_OVERRIDES.get(code.lower())
        base_code = override or fr_label or code
        desired_code = self._normalize_code(base_code)
        fallback_code = self._normalize_code(code)

        attr = Attribute.objects.filter(code=desired_code).first()
        if not attr:
            attr = Attribute.objects.filter(code=fallback_code).first()

        if not attr:
            target_code = self._unique_code(base_code)
            attr = Attribute.objects.create(code=target_code, data_type=data_type)
        else:
            if attr.data_type != data_type:
                self.stdout.write(
                    self.style.WARNING(
                        f"Attribute {attr.code} exists with data_type={attr.data_type}, expected {data_type}."
                    )
                )
            target_code = self._unique_code(base_code, current_code=attr.code)
            if attr.code != target_code:
                attr.code = target_code
                attr.save(update_fields=["code"])

        AttributeI18n.objects.get_or_create(
            attribute=attr,
            locale=locale,
            defaults={"label": fr_label or code},
        )
        ProductTypeAttribute.objects.get_or_create(product_type=product_type, attribute=attr, defaults={"required": False})
        return attr

    def _ensure_enum_value(self, attribute: Attribute, label: str, locale: Locale) -> AttributeValue:
        code = self._slugify(label)
        value, _ = AttributeValue.objects.get_or_create(attribute=attribute, code=code)
        AttributeValueI18n.objects.get_or_create(
            attribute_value=value,
            locale=locale,
            defaults={"label": label},
        )
        return value

    def _upsert_pav(
        self,
        product: Product,
        attribute: Attribute,
        *,
        attr_value: Optional[AttributeValue] = None,
        value_text: Optional[str] = None,
    ) -> ProductAttributeValue:
        pav, created = ProductAttributeValue.objects.get_or_create(
            product=product,
            attribute=attribute,
            defaults={
                "attribute_value": attr_value,
                "value_text": value_text,
            },
        )
        if not created:
            updates = []
            if pav.attribute_value_id != (attr_value.id if attr_value else None):
                pav.attribute_value = attr_value
                updates.append("attribute_value")
            if pav.value_text != value_text:
                pav.value_text = value_text
                updates.append("value_text")
            if updates:
                pav.save(update_fields=updates)
        return pav

    def _is_numeric(self, val: object) -> bool:
        try:
            float(str(val))
            return True
        except Exception:
            return False

    def _normalize_code(self, raw: str) -> str:
        slug = self._slugify(raw)
        if len(slug) <= 50:
            return slug
        digest = hashlib.md5(raw.encode("utf-8")).hexdigest()[:6]
        trimmed = slug[:43]
        return f"{trimmed}-{digest}"

    def _unique_code(self, base: str, current_code: Optional[str] = None) -> str:
        candidate = self._normalize_code(base)
        suffix = 2
        while Attribute.objects.filter(code=candidate).exclude(code=current_code).exists():
            candidate = self._normalize_code(f"{base}-{suffix}")
            suffix += 1
        return candidate
